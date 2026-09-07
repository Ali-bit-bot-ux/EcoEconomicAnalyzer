import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(10, 14, 23)        # #0A0E17 Deep cosmic navy
    CARD_BG = RGBColor(19, 27, 46)         # #131B2E Dark blue-gray card
    CARD_BORDER = RGBColor(37, 51, 77)     # #25334D Card border
    CARD_BG_ALT = RGBColor(24, 34, 58)     # #18223A
    ACCENT_ORANGE = RGBColor(249, 115, 22) # #F97316 Vibrant Orange
    ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 Satellite Cyan
    ACCENT_GREEN = RGBColor(16, 185, 129)  # #10B981 Emerald
    ACCENT_PURPLE = RGBColor(168, 85, 247) # #A855F7 Purple
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_LIGHT = RGBColor(226, 232, 240)   # #E2E8F0
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    def add_badge(slide, left, top, width, height, text, bg_color, text_color=TEXT_WHITE):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_color
        badge.line.fill.background()
        tf = badge.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        tf.margin_top = Inches(0.04)
        tf.margin_bottom = Inches(0.04)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.name = "Segoe UI"
        run.font.color.rgb = text_color
        return badge

    def add_header(slide, section_tag, main_title, subtitle=None):
        if section_tag:
            add_badge(slide, Inches(0.8), Inches(0.45), Inches(3.2), Inches(0.35), section_tag, CARD_BORDER, ACCENT_CYAN)
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = main_title
        run.font.bold = True
        run.font.size = Pt(24)
        run.font.name = "Segoe UI"
        run.font.color.rgb = TEXT_WHITE
        
        if subtitle:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(4)
            run2 = p2.add_run()
            run2.text = subtitle
            run2.font.size = Pt(13)
            run2.font.name = "Segoe UI"
            run2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: ТИТУЛЬНЫЙ (TITLE SLIDE)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    ring1 = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(-1.5), Inches(8.5), Inches(8.5))
    ring1.fill.background()
    ring1.line.color.rgb = RGBColor(30, 41, 69)
    ring1.line.width = Pt(1.5)

    ring2 = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-0.5), Inches(6.5), Inches(6.5))
    ring2.fill.background()
    ring2.line.color.rgb = RGBColor(40, 56, 92)
    ring2.line.width = Pt(1.0)

    add_badge(s1, Inches(1.0), Inches(0.9), Inches(5.2), Inches(0.42), 
              "СЕКЦИЯ: ЖЕР ЖӘНЕ ҒАРЫШ ТУРАЛЫ ҒЫЛЫМДАР. АСТРОНОМИЯ", RGBColor(19, 42, 69), ACCENT_CYAN)

    proj_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.0), Inches(1.0))
    tf = proj_box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "AgriCascade"
    run.font.bold = True
    run.font.size = Pt(50)
    run.font.name = "Segoe UI"
    run.font.color.rgb = ACCENT_ORANGE

    topic_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.7), Inches(10.5), Inches(1.5))
    tf2 = topic_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = 0
    p2 = tf2.paragraphs[0]
    run2 = p2.add_run()
    run2.text = "Ғарыштық радарлық-оптикалық зондтау және эконометрика негізінде агро-тәуекелдерді ерте болжау жүйесі"
    run2.font.bold = True
    run2.font.size = Pt(23)
    run2.font.name = "Segoe UI"
    run2.font.color.rgb = TEXT_WHITE

    p2_sub = tf2.add_paragraph()
    p2_sub.space_before = Pt(8)
    run2_sub = p2_sub.add_run()
    run2_sub.text = "Early Agro-Risk Forecasting System based on Space Radar-Optical Sensing and Econometrics"
    run2_sub.font.italic = True
    run2_sub.font.size = Pt(13)
    run2_sub.font.name = "Segoe UI"
    run2_sub.font.color.rgb = TEXT_MUTED

    card_authors = add_card(s1, Inches(1.0), Inches(4.7), Inches(5.3), Inches(1.9), CARD_BG, CARD_BORDER)
    tb_auth = s1.shapes.add_textbox(Inches(1.2), Inches(4.85), Inches(4.9), Inches(1.6))
    tf_a = tb_auth.text_frame
    tf_a.word_wrap = True
    p_at = tf_a.paragraphs[0]
    r = p_at.add_run()
    r.text = "ЖОБА ОРЫНДАУШЫЛАРЫ:"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT_CYAN
    
    p_a1 = tf_a.add_paragraph()
    p_a1.space_before = Pt(6)
    r = p_a1.add_run()
    r.text = "• Кәрімбай Әли (Karimbay Ali)\n• Төлебай Рамазан (Tolebay Ramazan)"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = TEXT_WHITE

    card_lead = add_card(s1, Inches(6.8), Inches(4.7), Inches(5.3), Inches(1.9), CARD_BG, CARD_BORDER)
    tb_lead = s1.shapes.add_textbox(Inches(7.0), Inches(4.85), Inches(4.9), Inches(1.6))
    tf_l = tb_lead.text_frame
    tf_l.word_wrap = True
    p_lt = tf_l.paragraphs[0]
    r = p_lt.add_run()
    r.text = "ҒЫЛЫМИ ЖЕТЕКШІ:"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT_ORANGE

    p_l1 = tf_l.add_paragraph()
    p_l1.space_before = Pt(6)
    r = p_l1.add_run()
    r.text = "• Искакова Айжан (Искакова А.)"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEXT_WHITE

    p_l2 = tf_l.add_paragraph()
    p_l2.space_before = Pt(4)
    r = p_l2.add_run()
    r.text = "Ғылыми кеңесші & Жоба кураторы"
    r.font.size = Pt(12)
    r.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: МӘСЕЛЕ (PROBLEMS)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "МӘСЕЛЕНІҢ ӨЗЕКТІЛІГІ", "Ғарыштық зондтаудың өзекті мәселелері", 
               "Дәстүрлі агро-мониторингтің шектеулері мен климаттық қауіп-қатерлер")

    card_data = [
        {
            "num": "01",
            "title": "Оптика шектеуі",
            "sub": "Бұлттылық және кеш тіркеу",
            "color": ACCENT_ORANGE,
            "text": "Дәстүрлі оптикалық спутниктер (Sentinel-2) бұлтты күндері жер бетін көре алмайды. Сонымен қатар, оптика өсімдіктің вегетациялық бұзылуын тек кеш кезеңде (жапырақтары сарғайып, зақымданғанда ғана) тіркейді."
        },
        {
            "num": "02",
            "title": "Тамыр тереңдігі",
            "sub": "NDVI тек жапырақ бетін көреді",
            "color": ACCENT_CYAN,
            "text": "Оптикалық спектрлік индекстер тек дақылдың үстіңгі қабатын сканерлейді. Топырақтың тамыр тереңдігіндегі (root-zone) ылғал тапшылығын көре алмайды, салдарынан фермерлер ерте әрекет ету мүмкіндігінен айырылады."
        },
        {
            "num": "03",
            "title": "Климаттық қауіп",
            "sub": "Аномалиялар мен баға шоктары",
            "color": RGBColor(244, 63, 94),
            "text": "Климаттың құбылуы мен экстремалды құрғақшылықты уақытылы анықтамау өңірлік астық көлемінің күрт төмендеуіне, логистикалық инфрақұрылым тоқырауына және биржада баға шоктарына алып келеді."
        }
    ]

    card_w = Inches(3.64)
    card_h = Inches(4.7)
    gap = Inches(0.39)
    start_x = Inches(0.8)
    start_y = Inches(1.85)

    for i, c in enumerate(card_data):
        cx = start_x + i * (card_w + gap)
        add_card(s2, cx, start_y, card_w, card_h)

        badge = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.3), start_y + Inches(0.35), Inches(0.8), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c["color"]
        badge.line.fill.background()
        btf = badge.text_frame
        btf.margin_top = Inches(0.04)
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = c["num"]
        br.font.bold = True
        br.font.size = Pt(13)
        br.font.color.rgb = TEXT_WHITE

        tb = s2.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(0.95), card_w - Inches(0.6), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        p_t = tf.paragraphs[0]
        r = p_t.add_run()
        r.text = c["title"]
        r.font.bold = True
        r.font.size = Pt(20)
        r.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.space_before = Pt(4)
        r = p_s.add_run()
        r.text = c["sub"]
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = c["color"]

        div = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), start_y + Inches(2.15), card_w - Inches(0.6), Pt(1.5))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

        tb_body = s2.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(2.3), card_w - Inches(0.6), Inches(2.1))
        tf_b = tb_body.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = 0
        p_b = tf_b.paragraphs[0]
        r = p_b.add_run()
        r.text = c["text"]
        r.font.size = Pt(13.5)
        r.font.name = "Segoe UI"
        r.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 3: KEY FEATURES
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "НЕГІЗГІ ЕРЕКШЕЛІКТЕР", "Ғарыштық мультиспектрлік синтез (Space Data Fusion)", 
               "Оптикалық, радарлық және микротолқынды зондтауды біріктіру арқылы ерте болжау")

    feat_data = [
        {
            "badge": "01 // РАДАРЛЫҚ ЗОНДТАУ",
            "title": "C-band Радиолокация",
            "sub": "Sentinel-1 (SAR) — ESA",
            "color": ACCENT_CYAN,
            "desc": "Радиолокациялық кері шашырау (backscatter) арқылы бұлттылық пен кез келген ауа райына қарамастан, тәуліктің кез келген мезгілінде топырақ беті мен өсімдік діңінің ылғалдылығын дәл тіркейді.",
            "stat": "24/7",
            "stat_label": "Ауа райына тәуелсіз бақылау"
        },
        {
            "badge": "02 // РАДИОМЕТРИЯ",
            "title": "L-band Радиометрия",
            "sub": "NASA SMAP Спутнигі",
            "color": ACCENT_PURPLE,
            "desc": "Микротолқынды радиометрия көмегімен жер асты тереңдігіне бойлап, өсімдік тамыр жүйесіндегі (Root-zone Soil Moisture) су қоры мен ылғал тапшылығы динамикасын тіркейді.",
            "stat": "0-100 см",
            "stat_label": "Тамыр тереңдігіндегі ылғал"
        },
        {
            "badge": "03 // НЕГІЗГІ БАСЫМДЫҚ",
            "title": "Ақпараттық басымдық",
            "sub": "Ерте ескерту жүйесі",
            "color": ACCENT_ORANGE,
            "desc": "Дәстүрлі оптикалық сарғаю (NDVI төмендеуі) көзге көрінгенге дейін тамырлық ылғал тапшылығын алдын ала анықтап, шешім қабылдауға құнды уақыт береді.",
            "stat": "+27 КҮН",
            "stat_label": "Оптикалық белгіден бұрын"
        }
    ]

    for i, f in enumerate(feat_data):
        cx = start_x + i * (card_w + gap)
        bg_col = CARD_BG if i < 2 else RGBColor(30, 25, 38)
        bord_col = CARD_BORDER if i < 2 else ACCENT_ORANGE
        add_card(s3, cx, start_y, card_w, card_h, bg_col, bord_col)

        add_badge(s3, cx + Inches(0.3), start_y + Inches(0.35), card_w - Inches(0.6), Inches(0.35), f["badge"], CARD_BG_ALT, f["color"])

        tb = s3.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(0.85), card_w - Inches(0.6), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        p_t = tf.paragraphs[0]
        r = p_t.add_run()
        r.text = f["title"]
        r.font.bold = True
        r.font.size = Pt(21)
        r.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.space_before = Pt(3)
        r = p_s.add_run()
        r.text = f["sub"]
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = f["color"]

        tb_body = s3.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(1.95), card_w - Inches(0.6), Inches(1.5))
        tf_b = tb_body.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = 0
        p_b = tf_b.paragraphs[0]
        r = p_b.add_run()
        r.text = f["desc"]
        r.font.size = Pt(13)
        r.font.color.rgb = TEXT_LIGHT

        stat_card = add_card(s3, cx + Inches(0.3), start_y + Inches(3.45), card_w - Inches(0.6), Inches(0.95), CARD_BG_ALT, CARD_BORDER)
        tb_st = s3.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(3.45), card_w - Inches(0.6), Inches(0.95))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        p_s1 = tf_st.paragraphs[0]
        p_s1.alignment = PP_ALIGN.CENTER
        r = p_s1.add_run()
        r.text = f["stat"]
        r.font.bold = True
        r.font.size = Pt(24 if i < 2 else 28)
        r.font.color.rgb = f["color"]

        p_s2 = tf_st.add_paragraph()
        p_s2.alignment = PP_ALIGN.CENTER
        r = p_s2.add_run()
        r.text = f["stat_label"]
        r.font.size = Pt(10.5)
        r.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: TECH STACK
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "ТЕХНОЛОГИЯЛЫҚ АРХИТЕКТУРА", "AgriCascade Tech Stack & Data Pipeline", 
               "Орбиталық деректерді өңдеуден бастап эконометрикалық болжау мен ГИС-интерфейске дейін")

    stack_modules = [
        {
            "num": "01",
            "name": "Orbital Data ETL",
            "tag": "ҒАРЫШТЫҚ ДЕРЕКТЕР ҚҰБЫРЫ",
            "color": ACCENT_CYAN,
            "techs": ["Google Earth Engine (GEE API)", "Copernicus Sentinel-1 (SAR)", "NASA SMAP Radiometer", "Sentinel-2 MSI (Optics)"],
            "desc": "Радарлық және микротолқынды ғарыштық деректерді жергілікті климаттық торға автоматты калибрлеу, сүзу және интеграциялау."
        },
        {
            "num": "02",
            "name": "Geo-Econometrics",
            "tag": "ЭКОНОМЕТРИКА ЖӘНЕ ТАЛДАУ",
            "color": ACCENT_ORANGE,
            "techs": ["Python 3.11+ / Pandas / NumPy", "Statsmodels & SciPy", "Vector Autoregression (VAR)", "Granger Causality Testing"],
            "desc": "Спутниктік аномалиялар мен баға шоктары арасындағы себеп-салдарлық байланысты дәлелдеу және бағаны модельдеу."
        },
        {
            "num": "03",
            "name": "GIS & Interface",
            "tag": "ИНТЕРАКТИВТІ КАРТОГРАФИЯ",
            "color": ACCENT_GREEN,
            "techs": ["Streamlit Web Framework", "Folium / Leaflet Geovisualization", "OSMnx Road & Railway Graphs", "207 Элеватор Бағыттары"],
            "desc": "Қазақстанның астық логистикасы мен элеватор инфрақұрылымын көрсететін интерактивті карталар мен басқару панелі."
        }
    ]

    for i, m in enumerate(stack_modules):
        cx = start_x + i * (card_w + gap)
        add_card(s4, cx, start_y, card_w, card_h)

        add_badge(s4, cx + Inches(0.3), start_y + Inches(0.35), card_w - Inches(0.6), Inches(0.32), m["tag"], CARD_BG_ALT, m["color"])

        tb = s4.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(0.8), card_w - Inches(0.6), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        p = tf.paragraphs[0]
        r_num = p.add_run()
        r_num.text = m["num"] + ". "
        r_num.font.bold = True
        r_num.font.size = Pt(20)
        r_num.font.color.rgb = m["color"]
        
        r_name = p.add_run()
        r_name.text = m["name"]
        r_name.font.bold = True
        r_name.font.size = Pt(20)
        r_name.font.color.rgb = TEXT_WHITE

        tb_techs = s4.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(1.75), card_w - Inches(0.6), Inches(1.8))
        tf_t = tb_techs.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = 0
        
        p_th = tf_t.paragraphs[0]
        r = p_th.add_run()
        r.text = "ҚОЛДАНЫЛҒАН ҚҰРАЛДАР:"
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = TEXT_MUTED

        for t in m["techs"]:
            pt = tf_t.add_paragraph()
            pt.space_before = Pt(3)
            r = pt.add_run()
            r.text = "✔ " + t
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = TEXT_WHITE

        div = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), start_y + Inches(3.6), card_w - Inches(0.6), Pt(1.5))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

        tb_desc = s4.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(3.75), card_w - Inches(0.6), Inches(0.85))
        tf_d = tb_desc.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = 0
        p_d = tf_d.paragraphs[0]
        r = p_d.add_run()
        r.text = m["desc"]
        r.font.size = Pt(12)
        r.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: FUNCTIONS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "ПЛАТФОРМА ФУНКЦИЯЛАРЫ", "Интерактивті ГИС-Дашборд модульдері (Live Platform)", 
               "Шешім қабылдауға арналған 4 негізгі талдау және мониторинг сервисі")

    funcs = [
        {
            "id": "01",
            "name": "PhenoMap",
            "subtitle": "Оптикалық вегетациялық мониторинг",
            "color": ACCENT_GREEN,
            "desc": "Sentinel-2 спутнигінің оптикалық спектрлік индекстерінің (NDVI/NDWI) геокеңістіктік динамикасы. Өсімдік жамылғысының биомассасы мен ылғал индексін нақты уақытта бақылау."
        },
        {
            "id": "02",
            "name": "HydroRisk",
            "subtitle": "Радарлық тамыр ылғалдылық картасы",
            "color": ACCENT_CYAN,
            "desc": "Sentinel-1 SAR және NASA SMAP радарлық-радиометриялық деректері негізінде топырақтың тамыр деңгейіндегі су балансы мен гидрологиялық күйзеліс аймақтарын анықтау."
        },
        {
            "id": "03",
            "name": "SiloSentry",
            "subtitle": "Элеваторлар инфрақұрылымы",
            "color": ACCENT_ORANGE,
            "desc": "Қазақстан бойынша 207 элеватордың кеңістіктік координаттары, логистикалық маршруттары (OSMnx) және астық сақтау сыйымдылығының кешенді желілік моделі."
        },
        {
            "id": "04",
            "name": "MarketSpread & RiskEngine",
            "subtitle": "Эконометрикалық ИИ-талдау",
            "color": ACCENT_PURPLE,
            "desc": "Орбиталық вегетациялық индикаторлар мен биржалық бағалар негізінде астық нарығындағы баға шоктарын және жеткізу тәуекелдерін алдын ала бағалау."
        }
    ]

    grid_w = Inches(5.66)
    grid_h = Inches(2.25)
    gap_x = Inches(0.4)
    gap_y = Inches(0.3)
    start_gx = Inches(0.8)
    start_gy = Inches(1.85)

    for i, fn in enumerate(funcs):
        row = i // 2
        col = i % 2
        fx = start_gx + col * (grid_w + gap_x)
        fy = start_gy + row * (grid_h + gap_y)

        add_card(s5, fx, fy, grid_w, grid_h)

        badge = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx + Inches(0.25), fy + Inches(0.25), Inches(0.65), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = fn["color"]
        badge.line.fill.background()
        btf = badge.text_frame
        btf.margin_top = Inches(0.04)
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = fn["id"]
        br.font.bold = True
        br.font.size = Pt(13)
        br.font.color.rgb = TEXT_WHITE

        tb_fn = s5.shapes.add_textbox(fx + Inches(1.05), fy + Inches(0.18), grid_w - Inches(1.3), Inches(0.7))
        tf_f = tb_fn.text_frame
        tf_f.word_wrap = True
        tf_f.margin_left = 0
        p = tf_f.paragraphs[0]
        r = p.add_run()
        r.text = fn["name"]
        r.font.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = TEXT_WHITE

        p_sub = tf_f.add_paragraph()
        r_s = p_sub.add_run()
        r_s.text = fn["subtitle"]
        r_s.font.size = Pt(11)
        r_s.font.bold = True
        r_s.font.color.rgb = fn["color"]

        tb_txt = s5.shapes.add_textbox(fx + Inches(0.25), fy + Inches(0.95), grid_w - Inches(0.5), Inches(1.15))
        tf_t = tb_txt.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = 0
        p_t = tf_t.paragraphs[0]
        r = p_t.add_run()
        r.text = fn["desc"]
        r.font.size = Pt(12.5)
        r.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 6: НАБЫҚ (MARKET OPPORTUNITY — TAM / SAM / SOM)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "НАРЫҚ КӨЛЕМІ", "Market Opportunity — TAM / SAM / SOM", 
               "Агро-геоаналитикалық бағдарламалық жасақтама мен ғарыштық мониторинг нарығы")

    markets = [
        {
            "abbr": "TAM",
            "title": "Total Addressable Market",
            "val": "$12.0B",
            "color": ACCENT_CYAN,
            "desc": "Жаһандық агро-геоаналитикалық БҚ және ғарыштық мониторингтік софт нарығының жалпы көлемі."
        },
        {
            "abbr": "SAM",
            "title": "Serviceable Addressable Market",
            "val": "$140M",
            "color": ACCENT_ORANGE,
            "desc": "Орталық Азия мен ТМД өңірінің ғарыштық агро-мониторинг және шешім қабылдау жүйелері нарығы."
        },
        {
            "abbr": "SOM",
            "title": "Serviceable Obtainable Market (3 жыл)",
            "val": "$1.2M",
            "color": ACCENT_GREEN,
            "desc": "Қазақстандағы 207 элеватор мен ірі астық холдингтерінің алғашқы 3 жылдағы мақсатты сегменті."
        }
    ]

    for i, m in enumerate(markets):
        cx = start_x + i * (card_w + gap)
        add_card(s6, cx, start_y, card_w, card_h)

        add_badge(s6, cx + Inches(0.3), start_y + Inches(0.35), Inches(1.1), Inches(0.4), m["abbr"], CARD_BG_ALT, m["color"])

        tb_val = s6.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(0.95), card_w - Inches(0.6), Inches(1.0))
        tf_v = tb_val.text_frame
        tf_v.word_wrap = True
        tf_v.margin_left = 0
        p_v = tf_v.paragraphs[0]
        r = p_v.add_run()
        r.text = m["val"]
        r.font.bold = True
        r.font.size = Pt(44)
        r.font.color.rgb = m["color"]

        p_sub = tf_v.add_paragraph()
        r_s = p_sub.add_run()
        r_s.text = m["title"]
        r_s.font.size = Pt(13)
        r_s.font.bold = True
        r_s.font.color.rgb = TEXT_WHITE

        div = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), start_y + Inches(2.35), card_w - Inches(0.6), Pt(1.5))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

        tb_desc = s6.shapes.add_textbox(cx + Inches(0.3), start_y + Inches(2.55), card_w - Inches(0.6), Inches(1.8))
        tf_d = tb_desc.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = 0
        p_d = tf_d.paragraphs[0]
        r = p_d.add_run()
        r.text = m["desc"]
        r.font.size = Pt(14)
        r.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 7: BACKTESTING & SCIENTIFIC VALIDATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "ҒЫЛЫМИ ВАЛИДАЦИЯ", "BackTesting: Орбиталық деректерді тексеру", 
               "Солтүстік Қазақстан облысындағы тарихи деректерге жүргізілген ретроспективті эксперимент")

    metrics = [
        {"val": "СҚО өңірі", "label": "Сынақ полигоны", "color": ACCENT_CYAN},
        {"val": "2021 жыл", "label": "Аномальды құрғақшылық", "color": ACCENT_ORANGE},
        {"val": "0%", "label": "Data Leakage (Ақпарат ақпауы)", "color": ACCENT_GREEN},
        {"val": "27 күн бұрын", "label": "Ерте анықтау нәтижесі", "color": ACCENT_PURPLE},
    ]

    metric_w = Inches(2.7)
    metric_gap = Inches(0.31)
    m_y = Inches(1.85)

    for i, m in enumerate(metrics):
        mx = start_x + i * (metric_w + metric_gap)
        add_card(s7, mx, m_y, metric_w, Inches(1.3), CARD_BG, CARD_BORDER)

        tb = s7.shapes.add_textbox(mx + Inches(0.15), m_y + Inches(0.15), metric_w - Inches(0.3), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = m["val"]
        r1.font.bold = True
        r1.font.size = Pt(21)
        r1.font.color.rgb = m["color"]

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = m["label"]
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    t_card = add_card(s7, start_x, Inches(3.45), Inches(11.73), Inches(3.1), CARD_BG, CARD_BORDER)

    tb_th = s7.shapes.add_textbox(start_x + Inches(0.4), Inches(3.6), Inches(10.9), Inches(0.4))
    tf_th = tb_th.text_frame
    p_th = tf_th.paragraphs[0]
    r = p_th.add_run()
    r.text = "РЕТРОСПЕКТИВТІ ЭКСПЕРИМЕНТ ТАЙМЛАЙНЫ (СҚО, 2021 ЖЫЛҒЫ МОДЕЛЬДЕУ)"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = ACCENT_ORANGE

    timeline_steps = [
        {
            "step": "Т - 27 КҮН",
            "title": "Тамыр ылғалының дағдарысы",
            "desc": "NASA SMAP L-band радиометриясы топырақтың тамыр тереңдігінде жасырын ылғал аномалиясын тіркеді.",
            "color": ACCENT_PURPLE
        },
        {
            "step": "Т - 14 КҮН",
            "title": "Өсімдік діңінің кебуі",
            "desc": "Sentinel-1 SAR радарлық кері шашырауы дің мен топырақ беткі қабатының күрт құрғауын растады.",
            "color": ACCENT_CYAN
        },
        {
            "step": "Т = 0 КҮН",
            "title": "Оптикалық сарғаю",
            "desc": "Sentinel-2 оптикалық суреттерінде дақылдардың жаппай сарғаюы (NDVI құлдырауы) бірінші рет көрінді.",
            "color": ACCENT_ORANGE
        },
        {
            "step": "Т + 10 КҮН",
            "title": "Ресми хабарлама & Баға шогы",
            "desc": "Құрғақшылық ресми түрде жарияланып, биржада астық бағасының күрт өсуі орын алды.",
            "color": RGBColor(244, 63, 94)
        }
    ]

    step_w = Inches(2.6)
    step_gap = Inches(0.24)
    step_x0 = start_x + Inches(0.4)
    step_y = Inches(4.15)

    for i, s in enumerate(timeline_steps):
        sx = step_x0 + i * (step_w + step_gap)
        
        badge = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, step_y, Inches(1.3), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = s["color"]
        badge.line.fill.background()
        btf = badge.text_frame
        btf.margin_top = Inches(0.03)
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = s["step"]
        br.font.bold = True
        br.font.size = Pt(11)
        br.font.color.rgb = TEXT_WHITE

        tb_step = s7.shapes.add_textbox(sx, step_y + Inches(0.45), step_w, Inches(1.6))
        tf_s = tb_step.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = 0
        p_st = tf_s.paragraphs[0]
        r = p_st.add_run()
        r.text = s["title"]
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = TEXT_WHITE

        p_sd = tf_s.add_paragraph()
        p_sd.space_before = Pt(4)
        r = p_sd.add_run()
        r.text = s["desc"]
        r.font.size = Pt(11.5)
        r.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 8: ҚОРЫТЫНДЫ (CONCLUSION & CONTACTS)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)

    ring_c = s8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(0.5), Inches(6.33), Inches(6.33))
    ring_c.fill.background()
    ring_c.line.color.rgb = RGBColor(25, 36, 62)
    ring_c.line.width = Pt(1.5)

    add_badge(s8, Inches(4.66), Inches(1.2), Inches(4.0), Inches(0.4), "AGRICASCADE PROJECT 2026", CARD_BORDER, ACCENT_ORANGE)

    tb_ty = s8.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(10.33), Inches(1.5))
    tf_ty = tb_ty.text_frame
    tf_ty.word_wrap = True
    p = tf_ty.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Назар аударғандарыңызға рақмет!"
    r.font.bold = True
    r.font.size = Pt(40)
    r.font.color.rgb = TEXT_WHITE

    p_sub = tf_ty.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(8)
    r = p_sub.add_run()
    r.text = "Ғарыштық технологиялар арқылы азық-түлік қауіпсіздігін ерте болжау"
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT_CYAN

    c_auth = add_card(s8, Inches(2.2), Inches(3.8), Inches(4.2), Inches(2.2), CARD_BG, CARD_BORDER)
    tb_ca = s8.shapes.add_textbox(Inches(2.4), Inches(4.0), Inches(3.8), Inches(1.8))
    tf_ca = tb_ca.text_frame
    tf_ca.word_wrap = True
    p = tf_ca.paragraphs[0]
    r = p.add_run()
    r.text = "АВТОРЛАР:"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT_CYAN

    p = tf_ca.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = "• Кәрімбай Әли\n• Төлебай Рамазан"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_WHITE

    p = tf_ca.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = "Жоба жасақтаушылары & Зерттеушілер"
    r.font.size = Pt(11)
    r.font.color.rgb = TEXT_MUTED

    c_sup = add_card(s8, Inches(6.9), Inches(3.8), Inches(4.2), Inches(2.2), CARD_BG, CARD_BORDER)
    tb_cs = s8.shapes.add_textbox(Inches(7.1), Inches(4.0), Inches(3.8), Inches(1.8))
    tf_cs = tb_cs.text_frame
    tf_cs.word_wrap = True
    p = tf_cs.paragraphs[0]
    r = p.add_run()
    r.text = "ҒЫЛЫМИ ЖЕТЕКШІ:"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT_ORANGE

    p = tf_cs.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = "• Искакова Айжан (Искакова А.)"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_WHITE

    p = tf_cs.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = "Ғылыми кеңесші & Жетекші"
    r.font.size = Pt(11)
    r.font.color.rgb = TEXT_MUTED

    # Save presentations
    output_workspace = r"d:\legion\code\Projects\daryn\AgriCascade_Presentation_Updated.pptx"
    output_parent = r"D:\legion\AgriCascade_Presentation_Updated.pptx"
    output_v2 = r"D:\legion\AgriCascade_Presentation_v2.pptx"

    try:
        prs.save(output_workspace)
        print(f"Saved successfully to: {output_workspace}")
    except Exception as e:
        print(f"Error saving to {output_workspace}: {e}")

    try:
        prs.save(output_parent)
        print(f"Saved successfully to: {output_parent}")
    except Exception as e:
        print(f"Could not overwrite {output_parent} directly (file is probably locked in PowerPoint: {e})")
        prs.save(output_v2)
        print(f"Saved alternative file to: {output_v2}")

if __name__ == "__main__":
    create_presentation()
